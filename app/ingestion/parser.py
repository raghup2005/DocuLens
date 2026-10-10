from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def parse_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""

    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n\n".join(pages)


def parse_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text.strip())

    return "\n".join(paragraphs)


def parse_txt(file_path: str) -> str:
    """Extract text from a TXT file."""

    path = Path(file_path)

    return path.read_text(
        encoding="utf-8",
        errors="ignore"
    )


def parse_pptx(file_path: str) -> str:
    """Extract text from a PowerPoint presentation."""

    presentation = Presentation(file_path)

    slides_text = []

    for slide in presentation.slides:

        slide_content = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_content.append(text)

        if slide_content:
            slides_text.append(
                "\n".join(slide_content)
            )

    return "\n\n".join(slides_text)

def parse_document(file_path: str, file_type: str) -> str:
    """
    Parse a document based on its detected file type.
    """

    if file_type == "pdf":
        return parse_pdf(file_path)

    elif file_type == "docx":
        return parse_docx(file_path)

    elif file_type == "text":
        return parse_txt(file_path)

    elif file_type == "powerpoint":
        return parse_pptx(file_path)

    else:
        raise ValueError(
            f"Unsupported document type for parser: {file_type}"
        )
def parse_pdf_by_page(file_path: str) -> list[dict]:
    """
    Extract PDF text while preserving page numbers.

    Returns:
        List of dictionaries containing page number and text.
    """

    reader = PdfReader(file_path)

    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        text = page.extract_text()

        if text and text.strip():

            pages.append({
                "page": page_number,
                "text": text.strip()
            })

    return pages
if __name__ == "__main__":
    file_path = "data/raw/1706.03762v7.pdf"

    text = parse_document(
        file_path,
        "pdf"
    )

    print("Current directory:", Path.cwd())
    print("File exists:", Path(file_path).exists())
    print("Absolute path:", Path(file_path).resolve())

    print(text[:2000])